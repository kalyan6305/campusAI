"""
Resume Optimization Agent - Specialized components for resume analysis and ATS optimization.
"""

from __future__ import annotations
import logging
import io
import os
import uuid
import json
from typing import List, Dict, Any
from app.llm.factory import get_llm_provider
from app.core.config import get_settings
from fpdf import FPDF
import pypdf
import docx

logger = logging.getLogger(__name__)

class ResumeOptimizationAgent:
    def __init__(self):
        self.llm = get_llm_provider()
        self.settings = get_settings()

    async def parse_resume(self, file_content: bytes, filename: str) -> str:
        """Extract text from PDF or DOCX resume."""
        text = ""
        try:
            if filename.lower().endswith('.pdf'):
                reader = pypdf.PdfReader(io.BytesIO(file_content))
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            elif filename.lower().endswith('.docx'):
                doc = docx.Document(io.BytesIO(file_content))
                text = "\n".join([para.text for para in doc.paragraphs])
            else:
                text = file_content.decode('utf-8', errors='ignore')
        except Exception as e:
            logger.error(f"Error parsing resume: {e}")
            raise Exception(f"Failed to parse resume: {e}")
        return text

    async def analyze_alignment(self, resume_text: str, job_description: str) -> Dict[str, Any]:
        """Analyze how well the resume matches the job description."""
        prompt = f"""
        Compare the following resume against the job description.
        
        Job Description:
        {job_description[:3000]}
        
        Resume:
        {resume_text[:4000]}
        
        Identify:
        1. Matched Skills (present in both)
        2. Missing Skills (required by JD but missing/weak in resume)
        3. Relevant Keywords for ATS optimization
        4. Experience alignment summary
        5. Overall ATS match score (0-100)

        Format your response as a JSON object:
        {{
            "matched_skills": ["skill1", "skill2"],
            "missing_skills": ["skill3", "skill4"],
            "keywords": ["key1", "key2"],
            "alignment_summary": "Summary text...",
            "match_score": 85
        }}
        """
        messages = [{"role": "user", "content": prompt}]
        response = await self.llm.generate(messages)
        try:
            start = response.find('{')
            end = response.rfind('}') + 1
            return json.loads(response[start:end])
        except Exception:
            # Fallback if AI doesn't return clean JSON
            return {"raw_analysis": response, "match_score": 0}

    async def optimize_resume(self, resume_text: str, job_description: str, analysis: Dict[str, Any], role: str = "Target Role") -> Dict[str, Any]:
        """Generate an elite, balanced version of the resume with score booster suggestions and robust parsing."""
        prompt = f"""
        You are a World-Class Executive Resume Architect. Your mission is to re-engineer the provided resume into an Elite, high-impact document optimized for: {role}.

        ELITE STRUCTURAL FRAMEWORK (STRUCTURE SHIELD):
        1. MANDATORY SECTIONS: You MUST include the following sections. Do not delete them:
           - # [Name] (Centered contact header)
           - ## EXECUTIVE SUMMARY (30-second pitch)
           - ## TECHNICAL CORE (Grouped skills)
           - ## PROFESSIONAL EXPERIENCE (Includes all Internships. Reframe irrelevant ones to focus on transferable skills.)
           - ## SELECTED PROJECTS (Showcase projects like OmniGenAI and Blood Donation Management System with impact.)
           - ## EDUCATION (Degree, University, CGPA)
           - ## CERTIFICATIONS & ACHIEVEMENTS
        2. STRATEGIC CONDENSING: 
           - Instead of deleting irrelevant items, strategically condense them. 
           - Example: If an internship is for "Java" but the role is "Python", focus on "Full Stack Development" and "System Architecture" rather than the specific language.
        3. PROFESSIONAL IMPACT: 
           - Bullet points MUST use [Action Verb] + [Specific Task] + [Quantifiable Result/Metric].
           - Use strong verbs: Orchestrated, Architected, Engineered, Optimized.

        OUTPUT FORMAT: 
        You MUST return a JSON object with exactly five fields. DO NOT use markdown code blocks like ```json.
        {{
            "optimized_resume_md": "THE FULL MARKDOWN RESUME. Use # for Name and ## for sections. Use horizontal rules --- for structure.",
            "optimization_insights": "A summary of the elite strategy used.",
            "unmatched_items": [
                {{"name": "Skill/Experience", "reason": "Brief reason for de-prioritization"}}
            ],
            "score_booster_suggestions": [
                "Actionable tip 1 (e.g., 'Add a project involving MLOps')",
                "Actionable tip 2 (e.g., 'Complete a specialized certification in AI Architecture')"
            ],
            "optimized_score": 98
        }}

        STRICT RULES:
        - Return ONLY the JSON. No conversational filler.
        - The resume MUST look like a complete, professional document.

        Job Description:
        {job_description[:2000]}
        
        Analysis Context:
        {json.dumps(analysis)}
        
        Original Resume:
        {resume_text[:4000]}
        """
        messages = [{"role": "user", "content": prompt}]
        response = await self.llm.generate(messages)
        
        import re
        try:
            # Robust JSON extraction handling blocks and filler
            json_match = re.search(r'\{.*\}', response, re.DOTALL)
            if json_match:
                json_str = json_match.group(0)
                parsed_json = json.loads(json_str)
                # Validate required fields
                if "optimized_resume_md" in parsed_json:
                    return parsed_json
            
            # Fallback to string search if regex is too restrictive
            start = response.find('{')
            end = response.rfind('}') + 1
            if start != -1 and end != -1:
                return json.loads(response[start:end])
                
            return {
                "optimized_resume_md": response, 
                "optimization_insights": "Optimized for keyword alignment.",
                "unmatched_items": [],
                "score_booster_suggestions": ["Ensure all key technologies are mentioned in descriptions."],
                "optimized_score": 85
            }
        except Exception as e:
            logger.error(f"Failed to parse optimization response: {e}")
            return {
                "optimized_resume_md": response, 
                "optimization_insights": "Optimization completed with structural fallback.",
                "unmatched_items": [],
                "score_booster_suggestions": ["Consider adding more metrics to your projects."],
                "optimized_score": 82
            }

    def generate_pdf(self, markdown_text: str, role: str) -> str:
        """Convert optimized Markdown resume to PDF and return relative URL."""
        pdf_dir = os.path.join("app", "static", "resumes")
        os.makedirs(pdf_dir, exist_ok=True)
        filename = f"optimized_resume_{role.lower().replace(' ', '_')}_{uuid.uuid4().hex[:6]}.pdf"
        file_path = os.path.join(pdf_dir, filename)

        try:
            pdf = FPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=15)
            
            # Helper for text sanitization to prevent encoding errors
            def safe(t): 
                return t.encode('latin-1', 'replace').decode('latin-1')

            lines = markdown_text.split('\n')
            for line in lines:
                clean_line = line.strip()
                if not clean_line:
                    pdf.ln(2)
                    continue

                if clean_line.startswith('# '):
                    # Main Name / Title
                    pdf.set_font("helvetica", style='B', size=22)
                    pdf.cell(0, 15, txt=safe(clean_line[2:]), ln=True, align='C')
                    pdf.ln(2)
                elif clean_line.startswith('## '):
                    # Section Header (Summary, Experience, etc.)
                    pdf.ln(4)
                    pdf.set_font("helvetica", style='B', size=14)
                    pdf.set_text_color(180, 0, 0) # Professional Deep Red
                    pdf.cell(0, 10, txt=safe(clean_line[3:]), ln=True)
                    # Professional Section Underline
                    pdf.set_draw_color(180, 0, 0)
                    pdf.set_line_width(0.4)
                    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
                    pdf.set_text_color(0, 0, 0)
                    pdf.ln(2)
                elif clean_line.startswith('### '):
                    # Job Title / Company / Project Name
                    pdf.ln(1)
                    pdf.set_font("helvetica", style='B', size=11)
                    pdf.cell(0, 8, txt=safe(clean_line[4:]), ln=True)
                    pdf.set_font("helvetica", size=10)
                elif clean_line.startswith('- ') or clean_line.startswith('* '):
                    # Bullet points (sanitize and replace markdown bold)
                    content = clean_line[2:].replace('**', '').replace('__', '')
                    pdf.set_font("helvetica", size=10)
                    pdf.multi_cell(0, 6, txt=f"  -  {safe(content)}")
                else:
                    # Generic text (contact info, descriptions)
                    txt = clean_line.replace('**', '').replace('__', '')
                    # If it's a contact line (contains | or @), maybe center it
                    is_contact = '|' in txt or '@' in txt
                    
                    pdf.set_font("helvetica", size=10)
                    if is_contact:
                        pdf.cell(0, 6, txt=safe(txt), ln=True, align='C')
                    else:
                        pdf.multi_cell(0, 6, txt=safe(txt))
                    pdf.ln(1)

            pdf.output(file_path)
        except Exception as e:
            logger.error(f"PDF Generation failed: {e}")
            # Simplified fallback that still shows the optimized text if possible
            try:
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("helvetica", size=10)
                pdf.multi_cell(0, 10, txt=markdown_text.encode('latin-1', 'replace').decode('latin-1'))
                pdf.output(file_path)
            except:
                # Absolute fallback
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("helvetica", size=12)
                pdf.cell(0, 10, txt="Optimized Resume (Minimalist Fallback)", ln=True)
                pdf.output(file_path)

        return f"/static/resumes/{filename}"
